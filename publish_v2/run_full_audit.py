"""publish_v2/run_full_audit.py - end-to-end audit pipeline.

Single sync, resumable, polite. No external API keys.
See run_full_audit.ps1 for CLI flags and usage.
"""
from __future__ import annotations

import argparse
import csv
import sys
import time
import traceback
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

from publish_v2 import discover  # noqa: E402
from publish_v2 import scrape_lite  # noqa: E402

CANONICAL = REPO / "publish_v2" / "data" / "canonical.csv"
V2_RESULTS = REPO / "scraper" / "data" / "scrape_results_enriched.csv"
PROGRESS = REPO / "publish_v2" / "data" / "coverage_progress.csv"
MANUAL_SEEDS = REPO / "publish_v2" / "data" / "manual_seeds.csv"

PROGRESS_FIELDS = [
    "slug", "name", "audit_status", "url", "discovery_method",
    "scrape_status", "indicators_evidence_found_count", "error", "attempted_at",
]


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def load_canonical() -> list[dict]:
    with CANONICAL.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def load_existing_v2() -> dict[str, dict]:
    if not V2_RESULTS.exists():
        return {}
    out = {}
    with V2_RESULTS.open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if r.get("scrape_status") == "ok":
                out[r["slug"]] = r
    return out


def load_excel_contacts() -> dict[str, dict]:
    excel = REPO / "Council_ClearSight_16.04.26.xlsx"
    if not excel.exists():
        return {}
    import openpyxl
    wb = openpyxl.load_workbook(excel, read_only=True, data_only=True)
    ws = wb["All Councils"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        d = dict(zip(header, row))
        slug = d.get("Slug")
        if not slug:
            continue
        out[slug] = {
            "url": (d.get("Website URL") or "").strip(),
            "email": (d.get("Email") or "").strip(),
            "phone": str(d.get("Phone") or "").strip(),
            "clerk_email": (d.get("Clerk Email") or "").strip(),
            "clerk_name": (d.get("Clerk Name") or "").strip(),
        }
    return out


def load_progress() -> dict[str, dict]:
    """Read coverage_progress.csv defensively. Skip rows missing 'slug'.
    Empty / corrupt / unreadable file -> start fresh."""
    out: dict[str, dict] = {}
    if not PROGRESS.exists() or PROGRESS.stat().st_size == 0:
        return out
    try:
        with PROGRESS.open(newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            if not reader.fieldnames or "slug" not in reader.fieldnames:
                print("  warn: coverage_progress.csv has no 'slug' column; "
                      "ignoring and starting fresh", file=sys.stderr)
                return out
            for r in reader:
                slug = (r.get("slug") or "").strip()
                if slug:
                    out[slug] = r
    except (csv.Error, UnicodeDecodeError, OSError) as e:
        print(f"  warn: coverage_progress.csv unreadable ({e}); "
              "ignoring and starting fresh", file=sys.stderr)
    return out


def append_progress(row: dict) -> None:
    """Append one row, retrying on Windows / OneDrive lock errors."""
    PROGRESS.parent.mkdir(parents=True, exist_ok=True)
    needs_header = (not PROGRESS.exists()) or PROGRESS.stat().st_size == 0
    last_err: Exception | None = None
    for attempt in range(5):
        try:
            with PROGRESS.open("a", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=PROGRESS_FIELDS)
                if needs_header:
                    w.writeheader()
                w.writerow({k: row.get(k, "") for k in PROGRESS_FIELDS})
            return
        except PermissionError as e:
            last_err = e
            time.sleep(0.5 + attempt)
    print(f"  warn: could not append to {PROGRESS.name} after 5 retries: "
          f"{last_err}", file=sys.stderr)


def append_enriched_row(row: dict, header: list[str]) -> None:
    new_file = not V2_RESULTS.exists()
    V2_RESULTS.parent.mkdir(parents=True, exist_ok=True)
    last_err: Exception | None = None
    for attempt in range(5):
        try:
            with V2_RESULTS.open("a", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=header, extrasaction="ignore")
                if new_file:
                    w.writeheader()
                w.writerow(row)
            return
        except PermissionError as e:
            last_err = e
            time.sleep(0.5 + attempt)
    raise last_err if last_err else RuntimeError("unknown append error")


def get_or_init_enriched_header(sample_row: dict) -> list[str]:
    if V2_RESULTS.exists() and V2_RESULTS.stat().st_size > 0:
        with V2_RESULTS.open(encoding="utf-8") as f:
            return next(csv.reader(f))
    return list(sample_row.keys())


def process_one(canonical_row, existing_v2, excel_contacts, manual_seeds,
                disc_session, gate, scrape_session, enriched_header_holder,
                discover_only=False, no_discover=False):
    slug = canonical_row["slug"]
    name = canonical_row["name"]
    out = {
        "slug": slug, "name": name, "audit_status": "pending",
        "url": "", "discovery_method": "", "scrape_status": "",
        "indicators_evidence_found_count": "", "error": "",
        "attempted_at": now_iso(),
    }

    if slug in existing_v2:
        prev = existing_v2[slug]
        out["audit_status"] = "verified"
        out["url"] = prev.get("website_url", "")
        out["discovery_method"] = "pre-existing"
        out["scrape_status"] = "ok"
        out["indicators_evidence_found_count"] = prev.get(
            "indicators_evidence_found_count", "")
        return out

    excel = excel_contacts.get(slug, {})
    url = excel.get("url", "")
    method = "excel" if url else ""

    if not url and not no_discover:
        try:
            res = discover.discover(slug, name, disc_session,
                                    manual_seeds=manual_seeds)
        except Exception as e:
            out["audit_status"] = "discovery_error"
            out["error"] = f"{type(e).__name__}: {e}"
            return out
        if res:
            url = res.url
            method = res.method

    if not url:
        out["audit_status"] = "no_web_presence_found"
        return out

    out["url"] = url
    out["discovery_method"] = method or "unknown"

    if discover_only:
        out["audit_status"] = "url_discovered_pending_scrape"
        return out

    try:
        scraped = scrape_lite.scrape_council(
            slug=slug, name=name, url=url,
            council_id=canonical_row.get("council_id") or "",
            contact_email=excel.get("email", ""),
            clerk_email=excel.get("clerk_email", ""),
            phone=excel.get("phone", ""),
            gate=gate, session=scrape_session,
        )
    except Exception as e:
        traceback.print_exc()
        out["audit_status"] = "scrape_error"
        out["error"] = f"{type(e).__name__}: {e}"
        return out

    out["scrape_status"] = scraped.get("scrape_status") or ""
    out["indicators_evidence_found_count"] = scraped.get(
        "indicators_evidence_found_count") or ""
    out["error"] = scraped.get("error") or ""

    try:
        if not enriched_header_holder:
            enriched_header_holder.extend(get_or_init_enriched_header(scraped))
        append_enriched_row(scraped, enriched_header_holder)
    except Exception as e:
        print(f"  warn: enriched-CSV append failed for {slug}: {e}",
              file=sys.stderr)

    if out["scrape_status"] == "ok":
        out["audit_status"] = "verified"
    elif out["scrape_status"] == "robots-disallow":
        out["audit_status"] = "blocked_by_robots"
    elif "unreach" in (out["scrape_status"] or "").lower():
        out["audit_status"] = "url_unreachable"
    else:
        out["audit_status"] = "scrape_failed"
    return out


def main():
    p = argparse.ArgumentParser(
        description="End-to-end Council ClearSight audit pipeline")
    p.add_argument("--limit", type=int, default=0)
    p.add_argument("--rescan", action="store_true")
    p.add_argument("--discover-only", action="store_true")
    p.add_argument("--no-discover", action="store_true")
    p.add_argument("--skip-publish", action="store_true")
    p.add_argument("--filter-source",
                   choices=["excel", "ons_only"], default=None)
    args = p.parse_args()

    print("=" * 60)
    print(" Council ClearSight - end-to-end audit pipeline")
    print("=" * 60)
    print(f" Started: {now_iso()}")
    print()

    canonical = load_canonical()
    print(f"Loaded canonical list: {len(canonical)} councils")
    if args.filter_source:
        canonical = [r for r in canonical
                     if r.get("source") == args.filter_source]
        print(f"Filtered to source={args.filter_source}: {len(canonical)}")

    existing_v2 = load_existing_v2()
    print(f"Loaded existing v2 results (ok status): {len(existing_v2)}")

    excel_contacts = load_excel_contacts()
    print(f"Loaded Excel contacts: {len(excel_contacts)}")

    manual_seeds = discover.load_manual_seeds(MANUAL_SEEDS)
    print(f"Loaded manual seeds: {len(manual_seeds)}")

    progress = {} if args.rescan else load_progress()
    if progress:
        print(f"Resuming - {len(progress)} councils already processed")
    if args.rescan and PROGRESS.exists():
        try:
            PROGRESS.unlink()
        except PermissionError:
            print("  warn: cannot unlink locked progress file; "
                  "will reuse it", file=sys.stderr)

    disc_session = discover.make_session()
    scrape_session = scrape_lite.make_session()
    gate = scrape_lite.PoliteGate()
    enriched_header_holder: list[str] = []

    todo = [r for r in canonical if r["slug"] not in progress]
    if args.limit:
        todo = todo[:args.limit]
    print(f"To process: {len(todo)}")
    print()

    counts: dict = defaultdict(int)
    started = time.time()

    for i, row in enumerate(todo, 1):
        try:
            result = process_one(
                row, existing_v2, excel_contacts, manual_seeds,
                disc_session, gate, scrape_session, enriched_header_holder,
                discover_only=args.discover_only,
                no_discover=args.no_discover,
            )
        except KeyboardInterrupt:
            print("\nInterrupted by user. Progress saved.")
            return 130
        except Exception:
            traceback.print_exc()
            result = {
                "slug": row["slug"], "name": row["name"],
                "audit_status": "fatal_error", "url": "",
                "discovery_method": "", "scrape_status": "",
                "indicators_evidence_found_count": "",
                "error": "uncaught: see stderr", "attempted_at": now_iso(),
            }
        counts[result["audit_status"]] += 1
        append_progress(result)

        if i % 25 == 0 or i == len(todo):
            elapsed = time.time() - started
            rate = i / elapsed if elapsed > 0 else 0
            eta = (len(todo) - i) / rate if rate > 0 else 0
            verified = counts.get("verified", 0)
            print(f"  [{i:>5}/{len(todo)}] verified={verified}  "
                  f"rate={rate:.2f}/s  eta={eta/60:.1f}m  "
                  f"counts={dict(counts)}")

    print()
    print("=" * 60)
    print(" Audit run complete")
    print("=" * 60)
    print(f" Finished: {now_iso()}")
    print(" Counts:")
    for status, n in sorted(counts.items(), key=lambda x: -x[1]):
        print(f"   {status:30s} {n:>6}")
    print()

    if not args.skip_publish:
        print("Running publish_v2/publish.py to regenerate per-slug JSONs...")
        import subprocess
        try:
            subprocess.run(
                [sys.executable, str(REPO / "publish_v2" / "publish.py")],
                check=True, cwd=REPO)
        except subprocess.CalledProcessError as e:
            print(f"  publish failed: {e}")
            return 1
        print("Publish complete.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
